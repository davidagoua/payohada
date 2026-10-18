"""Tests du module « Éditions » : états de paie collectifs.

Deux exigences structurent ces tests :

* **cohérence des chiffres** — un état ne recalcule rien, il agrège les lignes
  du bulletin ; les totaux doivent donc se recouper d'un état à l'autre ;
* **cloisonnement** — un état collectif expose toutes les rémunérations d'un
  dossier : il est réservé au cabinet.
"""
import unittest
from io import BytesIO

from tests import base
from tests.base import (
    API, SessionLocal, auth_headers, client, make_contrat, make_etablissement,
    make_salarie, make_tenant, make_user, reset_database,
)

from app.models import models as M
from app.services import editions as service
import app.services.editions_paie  # noqa: F401  (enregistre le catalogue)
from app.services.payroll import calculate_payslip


def _preparer_paie(db, t, nombre_salaries=3, mois=(5, 6), annee=2025, brut_base=300_000.0):
    """Crée des salariés, leurs bulletins calculés et une prime mensuelle."""
    contrats = [t["contrat"]]
    for index in range(1, nombre_salaries):
        salarie = make_salarie(
            db, t["etablissement"], matricule=f"E{index:03d}",
            nom=f"NOM{index}", prenom=f"Prenom{index}",
        )
        salarie.numero_securite_sociale = f"1234567890{index:03d}"
        contrats.append(make_contrat(
            db, t["dossier"], t["etablissement"], salarie,
            numero=f"C{index:03d}", salaire_mensuel=brut_base + index * 100_000,
        ))
    db.commit()

    for contrat in contrats:
        for mois_paie in mois:
            db.add(M.Prime(
                contrat_id=contrat.id, code="PRIME_REND",
                libelle="Prime de rendement", montant=30_000.0,
                mois=mois_paie, annee=str(annee),
            ))
            bulletin = calculate_payslip(db, contrat.id, mois_paie, annee)
            bulletin.statut = "calcule"
    db.commit()
    return contrats


class CatalogueEditionsTests(unittest.TestCase):
    def test_les_cinq_etats_sont_enregistres(self):
        codes = {d.code for d in service.catalogue()}
        self.assertEqual(
            codes,
            {"JOURNAL_PAIE", "LIVRE_PAIE", "CNPS_COTISATIONS", "ITS", "CHARGES_PATRONALES"},
        )

    def test_regroupement_par_famille(self):
        familles = service.catalogue_par_famille()
        self.assertIn(service.FAMILLE_PAIE, familles)
        self.assertEqual(len(familles[service.FAMILLE_PAIE]), 5)

    def test_chaque_etat_declare_ses_formats(self):
        for definition in service.catalogue():
            with self.subTest(code=definition.code):
                self.assertIn("pdf", definition.formats)
                self.assertTrue(definition.description)


class GenerationEditionsTests(unittest.TestCase):
    """Cohérence des chiffres produits par chaque état."""

    def setUp(self):
        reset_database()
        self.db = SessionLocal()
        self.t = make_tenant(self.db, "cabinet-editions")
        _preparer_paie(self.db, self.t)
        self.contexte = service.ContexteEdition(
            etablissement_id=self.t["etablissement"].id,
            mois_debut=5, annee_debut=2025, mois_fin=6, annee_fin=2025,
        )

    def tearDown(self):
        self.db.close()

    def _etat(self, code):
        return service.generer_edition(self.db, code, self.contexte)

    def test_journal_couvre_tous_les_bulletins(self):
        journal = self._etat("JOURNAL_PAIE")
        # 3 salariés × 2 mois
        self.assertEqual(journal.nombre_lignes, 6)
        self.assertIn("brut", journal.totaux)
        self.assertGreater(journal.totaux["brut"], 0)
        self.assertAlmostEqual(
            journal.totaux["brut"],
            sum(ligne["brut"] for ligne in journal.lignes),
            places=2,
        )

    def test_le_livre_detaille_les_rubriques(self):
        livre = self._etat("LIVRE_PAIE")
        journal = self._etat("JOURNAL_PAIE")
        # Le livre compte bien plus de lignes que le journal
        self.assertGreater(livre.nombre_lignes, journal.nombre_lignes * 5)
        # Les gains du livre égalent le brut du journal
        self.assertAlmostEqual(
            livre.totaux["gain"], journal.totaux["brut"], places=2
        )
        # Les retenues du livre égalent les cotisations salariales du journal
        self.assertAlmostEqual(
            livre.totaux["retenue"], journal.totaux["cotisations"], places=2
        )

    def test_its_net_egale_impot_brut_moins_ricf(self):
        etat = self._etat("ITS")
        self.assertEqual(etat.nombre_lignes, 6)
        self.assertAlmostEqual(
            etat.totaux["its_net"],
            round(etat.totaux["impot_brut"] - etat.totaux["ricf"], 2),
            places=2,
        )
        for ligne in etat.lignes:
            with self.subTest(matricule=ligne["matricule"]):
                self.assertAlmostEqual(
                    ligne["its_net"], round(ligne["impot_brut"] - ligne["ricf"], 2),
                    places=2,
                )

    def test_le_journal_reprend_l_its_de_l_etat_its(self):
        """L'ITS du journal de paie doit égaler celui de l'état dédié."""
        journal = self._etat("JOURNAL_PAIE")
        etat_its = self._etat("ITS")
        self.assertAlmostEqual(
            journal.totaux["its"], etat_its.totaux["its_net"], places=2
        )

    def test_cotisations_cnps_conformes_aux_taux(self):
        etat = self._etat("CNPS_COTISATIONS")
        self.assertEqual(etat.nombre_lignes, 6)
        # Retraite salariale : 6,3 % de la base CNPS
        self.assertAlmostEqual(
            etat.totaux["retraite_s"], round(etat.totaux["base"] * 0.063, 2), places=2
        )
        # Retraite patronale : 7,7 % de la base
        self.assertAlmostEqual(
            etat.totaux["retraite_p"], round(etat.totaux["base"] * 0.077, 2), places=2
        )
        # Total patronal = somme des quatre branches
        for ligne in etat.lignes:
            with self.subTest(matricule=ligne["matricule"]):
                self.assertAlmostEqual(
                    ligne["total_patronal"],
                    round(ligne["retraite_p"] + ligne["pf"] + ligne["at"] + ligne["maternite"], 2),
                    places=2,
                )

    def test_charges_patronales_totalisees(self):
        etat = self._etat("CHARGES_PATRONALES")
        self.assertEqual(etat.nombre_lignes, 6)
        for ligne in etat.lignes:
            with self.subTest(matricule=ligne["matricule"]):
                attendu = sum(
                    ligne[cle] for cle in
                    ("retraite", "pf", "at", "maternite", "cmu", "cn", "ta", "tfc")
                )
                self.assertAlmostEqual(ligne["total"], round(attendu, 2), places=2)
        self.assertAlmostEqual(
            etat.totaux["total"],
            sum(ligne["total"] for ligne in etat.lignes),
            places=2,
        )

    def test_les_brouillons_sont_exclus(self):
        """Un état de paie ne doit refléter que des bulletins arrêtés."""
        bulletin = self.db.query(M.BulletinPaie).first()
        bulletin.statut = "brouillon"
        self.db.commit()
        journal = self._etat("JOURNAL_PAIE")
        self.assertEqual(journal.nombre_lignes, 5)

    def test_periode_sans_donnees_produit_un_etat_vide_explicite(self):
        contexte = service.ContexteEdition(
            etablissement_id=self.t["etablissement"].id,
            mois_debut=1, annee_debut=2030, mois_fin=1, annee_fin=2030,
        )
        journal = service.generer_edition(self.db, "JOURNAL_PAIE", contexte)
        self.assertEqual(journal.nombre_lignes, 0)
        self.assertTrue(journal.notes, "une période vide doit être signalée")

    def test_isolement_entre_etablissements_du_meme_dossier(self):
        """Un état ne couvre que son établissement, pas le dossier entier.

        C'est la propriété centrale du périmètre : deux établissements d'un même
        dossier ont chacun leurs salariés et leurs bulletins.
        """
        autre_etab = make_etablissement(self.db, self.t["dossier"], code="ETAB2")
        salarie = make_salarie(self.db, autre_etab, matricule="AUTRE1",
                               nom="AUTRE", prenom="Salarie")
        contrat = make_contrat(self.db, self.t["dossier"], autre_etab, salarie,
                               numero="CAUTRE", salaire_mensuel=900_000.0)
        self.db.commit()
        bulletin = calculate_payslip(self.db, contrat.id, 6, 2025)
        bulletin.statut = "calcule"
        self.db.commit()

        # L'établissement d'origine ne voit pas ce salarié
        journal = self._etat("JOURNAL_PAIE")
        self.assertEqual(journal.nombre_lignes, 6)
        self.assertNotIn("AUTRE1", {ligne["matricule"] for ligne in journal.lignes})

        # Le second établissement ne voit que le sien, et son propre brut
        contexte_autre = service.ContexteEdition(
            etablissement_id=autre_etab.id,
            mois_debut=6, annee_debut=2025, mois_fin=6, annee_fin=2025,
        )
        journal_autre = service.generer_edition(self.db, "JOURNAL_PAIE", contexte_autre)
        self.assertEqual(journal_autre.nombre_lignes, 1)
        self.assertEqual(journal_autre.lignes[0]["matricule"], "AUTRE1")
        self.assertAlmostEqual(journal_autre.totaux["brut"], 900_000.0, places=2)

    def test_l_en_tete_nomme_l_etablissement(self):
        contexte = service.ContexteEdition(
            etablissement_id=self.t["etablissement"].id,
            mois_debut=6, annee_debut=2025, mois_fin=6, annee_fin=2025,
        )
        journal = service.generer_edition(self.db, "JOURNAL_PAIE", contexte)
        self.assertTrue(journal.contexte["etablissement"])
        self.assertIn("societe", journal.contexte)


class RendusEditionsTests(unittest.TestCase):
    """Les trois sorties doivent être des fichiers valides."""

    def setUp(self):
        reset_database()
        self.db = SessionLocal()
        self.t = make_tenant(self.db, "cabinet-editions-rendus")
        _preparer_paie(self.db, self.t, nombre_salaries=2, mois=(6,))
        self.contexte = service.ContexteEdition(
            etablissement_id=self.t["etablissement"].id,
            mois_debut=6, annee_debut=2025, mois_fin=6, annee_fin=2025,
        )

    def tearDown(self):
        self.db.close()

    def test_les_trois_formats_pour_chaque_etat(self):
        for definition in service.catalogue():
            resultat = service.generer_edition(self.db, definition.code, self.contexte)
            with self.subTest(code=definition.code):
                pdf = service.rendre(resultat, "pdf", paysage=definition.paysage)
                self.assertTrue(pdf.startswith(b"%PDF-"))
                self.assertTrue(pdf.rstrip().endswith(b"%%EOF"))

                excel = service.rendre(resultat, "excel")
                self.assertTrue(excel.startswith(b"PK"))  # archive openpyxl

                csv = service.rendre(resultat, "csv")
                self.assertTrue(csv.startswith(b"\xef\xbb\xbf"))  # BOM UTF-8

    def test_le_pdf_contient_les_donnees_de_l_etat(self):
        from pypdf import PdfReader

        resultat = service.generer_edition(self.db, "JOURNAL_PAIE", self.contexte)
        pdf = service.rendre(resultat, "pdf", paysage=True)
        texte = "".join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)

        # Le gabarit met le titre en capitales : on compare sans tenir compte
        # de la casse, comme le ferait un lecteur.
        self.assertIn("journal de paie", texte.lower())
        # Les totaux doivent se retrouver dans le document
        self.assertIn(service.format_montant(resultat.totaux["brut"]), texte)
        premier = resultat.lignes[0]
        self.assertIn(premier["matricule"], texte)

    def test_l_excel_contient_les_donnees(self):
        from openpyxl import load_workbook

        resultat = service.generer_edition(self.db, "ITS", self.contexte)
        contenu = service.rendre(resultat, "excel")
        feuille = load_workbook(BytesIO(contenu)).active

        valeurs = [
            cellule.value
            for ligne in feuille.iter_rows() for cellule in ligne
            if cellule.value is not None
        ]
        self.assertIn("Brut imposable", valeurs)
        self.assertIn(resultat.lignes[0]["matricule"], valeurs)

    def test_nom_de_fichier_exploitable(self):
        resultat = service.generer_edition(self.db, "JOURNAL_PAIE", self.contexte)
        nom = service.nom_fichier(resultat, "pdf")
        self.assertTrue(nom.endswith(".pdf"))
        self.assertIn("JOURNAL_PAIE", nom)
        for interdit in (" ", "/", "\r", "\n", '"'):
            self.assertNotIn(interdit, nom)

    def test_format_inconnu_refuse(self):
        resultat = service.generer_edition(self.db, "JOURNAL_PAIE", self.contexte)
        with self.assertRaises(ValueError):
            service.rendre(resultat, "docx")

    def test_le_csv_est_relisible_par_un_tableur(self):
        """Les montants sortent en nombres, jamais en flottants de Python."""
        import csv as module_csv
        import io

        resultat = service.generer_edition(self.db, "JOURNAL_PAIE", self.contexte)
        contenu = service.rendre(resultat, "csv").decode("utf-8-sig")
        lignes = list(module_csv.reader(io.StringIO(contenu), delimiter=";"))
        entetes, corps, totaux = lignes[0], lignes[1:-1], lignes[-1]
        self.assertEqual(entetes[0], resultat.colonnes[0].libelle)
        self.assertEqual(len(corps), resultat.nombre_lignes)

        positions = {c.cle: index for index, c in enumerate(resultat.colonnes)}
        montants = [c for c in resultat.colonnes if c.type == service.TYPE_MONTANT]
        self.assertTrue(montants)
        for ligne, modele in zip(corps, resultat.lignes):
            for colonne in montants:
                cellule = ligne[positions[colonne.cle]]
                self.assertNotIn(".", cellule, f"{colonne.cle} sort en flottant : {cellule}")
                self.assertEqual(int(cellule), round(modele[colonne.cle]))

        # Le pied d'état est libellé et ses montants restent des nombres
        self.assertEqual(totaux[0], service.LIBELLE_TOTAUX)
        self.assertEqual(
            int(totaux[positions["brut"]]), round(resultat.totaux["brut"])
        )

    def test_la_ligne_de_totaux_est_libellee_dans_tous_les_formats(self):
        """L'aperçu et les fichiers nomment la ligne de totaux de la même façon."""
        from openpyxl import load_workbook
        from pypdf import PdfReader

        resultat = service.generer_edition(self.db, "JOURNAL_PAIE", self.contexte)
        self.assertTrue(resultat.totaux)
        attendu = service.LIBELLE_TOTAUX

        pdf = service.rendre(resultat, "pdf", paysage=True)
        texte = "".join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)
        self.assertIn(attendu, texte)

        feuille = load_workbook(BytesIO(service.rendre(resultat, "excel"))).active
        valeurs = [
            cellule.value
            for ligne in feuille.iter_rows() for cellule in ligne
            if cellule.value is not None
        ]
        self.assertIn(attendu, valeurs)

        csv = service.rendre(resultat, "csv").decode("utf-8-sig")
        self.assertIn(attendu, csv.splitlines()[-1])

    def test_le_bandeau_ne_repete_pas_la_societe(self):
        """Le titre est déjà la société : elle ne doit pas suivre en dessous."""
        from pypdf import PdfReader

        resultat = service.generer_edition(self.db, "JOURNAL_PAIE", self.contexte)
        societe = resultat.contexte["societe"]
        pdf = service.rendre(resultat, "pdf", paysage=True)
        texte = "".join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)
        # Le titre est composé en capitales par la feuille de style.
        self.assertEqual(texte.lower().count(societe.lower()), 1)


class PeriodeEditionTests(unittest.TestCase):
    def test_periodes_enumerees_dans_l_ordre(self):
        contexte = service.ContexteEdition(
            etablissement_id=1, mois_debut=11, annee_debut=2024, mois_fin=2, annee_fin=2025,
        )
        self.assertEqual(
            contexte.periodes,
            [(2024, 11), (2024, 12), (2025, 1), (2025, 2)],
        )

    def test_libelle_simplifie_si_un_seul_mois(self):
        contexte = service.ContexteEdition(
            etablissement_id=1, mois_debut=6, annee_debut=2025, mois_fin=6, annee_fin=2025,
        )
        self.assertEqual(contexte.libelle_periode, "Juin 2025")

    def test_libelle_etendue_si_plusieurs_mois(self):
        contexte = service.ContexteEdition(
            etablissement_id=1, mois_debut=1, annee_debut=2025, mois_fin=3, annee_fin=2025,
        )
        self.assertEqual(contexte.libelle_periode, "Janvier 2025 – Mars 2025")

    def test_mois_invalide_refuse(self):
        with self.assertRaises(ValueError):
            service.ContexteEdition(
                etablissement_id=1, mois_debut=13, annee_debut=2025, mois_fin=1, annee_fin=2025,
            )


class FormatCsvTests(unittest.TestCase):
    """Le CSV est relu par un tableur, pas par un lecteur d'écran."""

    colonne_montant = service.Colonne("brut", "Brut", service.TYPE_MONTANT)
    colonne_nombre = service.Colonne("parts", "Parts", service.TYPE_NOMBRE)

    def test_montant_entier_sans_separateur_de_milliers(self):
        self.assertEqual(service.format_csv(1234567.0, self.colonne_montant), "1234567")

    def test_montant_decimal_a_la_virgule(self):
        self.assertEqual(service.format_csv(250000.5, self.colonne_montant), "250000,50")

    def test_montant_absent_reste_vide(self):
        self.assertEqual(service.format_csv(None, self.colonne_montant), "")

    def test_nombre_decimal_a_la_virgule(self):
        """Les nombres gardent la présentation française du reste de l'application."""
        self.assertEqual(service.format_csv(2.5, self.colonne_nombre), "2,50")

    def test_residu_de_flottant_est_arrondi(self):
        """Un montant calculé ne doit pas traîner ses décimales binaires."""
        self.assertEqual(
            service.format_csv(250000.49999999999, self.colonne_montant), "250000,50"
        )
        self.assertEqual(
            service.format_csv(1166.3999999999999, self.colonne_montant), "1166,40"
        )

    def test_libelle_des_totaux_porte_sur_la_premiere_colonne_de_texte(self):
        colonnes = [
            service.Colonne("matricule", "Matricule", service.TYPE_TEXTE),
            service.Colonne("brut", "Brut", service.TYPE_MONTANT, total=True),
        ]
        self.assertEqual(service.index_libelle_totaux(colonnes), 0)

    def test_libelle_des_totaux_absent_si_tout_est_totalise(self):
        colonnes = [service.Colonne("brut", "Brut", service.TYPE_MONTANT, total=True)]
        self.assertIsNone(service.index_libelle_totaux(colonnes))


class EditionsApiTests(unittest.TestCase):
    def setUp(self):
        reset_database()
        self.db = SessionLocal()
        self.t = make_tenant(self.db, "cabinet-editions-api")
        _preparer_paie(self.db, self.t, nombre_salaries=2, mois=(6,))
        self.entetes = auth_headers(self.t["user"])
        self.demande = {
            "etablissement_id": self.t["etablissement"].id,
            "mois_debut": 6, "annee_debut": 2025,
            "mois_fin": 6, "annee_fin": 2025,
        }

    def tearDown(self):
        self.db.close()

    def test_catalogue(self):
        r = client.get(f"{API}/editions", headers=self.entetes)
        self.assertEqual(r.status_code, 200, r.text[:200])
        familles = r.json()
        total = sum(len(f["editions"]) for f in familles)
        self.assertEqual(total, 5)

    def test_apercu_d_un_etat(self):
        r = client.post(f"{API}/editions/JOURNAL_PAIE/generer",
                        headers=self.entetes, json=self.demande)
        self.assertEqual(r.status_code, 200, r.text[:300])
        corps = r.json()
        self.assertEqual(corps["nombre_lignes"], 2)
        self.assertTrue(corps["colonnes"])
        self.assertIn("brut", corps["totaux"])
        self.assertIn("societe", corps["contexte"])

    def test_export_dans_les_trois_formats(self):
        attendus = {
            "pdf": b"%PDF-",
            "excel": b"PK",
            "csv": b"\xef\xbb\xbf",
        }
        for format_sortie, entete in attendus.items():
            with self.subTest(format=format_sortie):
                r = client.post(
                    f"{API}/editions/JOURNAL_PAIE/export",
                    headers=self.entetes, json=self.demande,
                    params={"format": format_sortie},
                )
                self.assertEqual(r.status_code, 200, r.text[:200])
                self.assertTrue(r.content.startswith(entete))
                self.assertIn("attachment", r.headers["content-disposition"])
                self.assertEqual(r.headers["x-nombre-lignes"], "2")

    def test_envoi_par_email_joint_le_pdf(self):
        appels = []

        def faux_send_email(to_email, subject, html, pieces_jointes=None):
            appels.append((to_email, subject, list(pieces_jointes or [])))
            return True, "ok"

        import app.routers.editions as routeur
        originel = routeur.send_email
        routeur.send_email = faux_send_email
        try:
            r = client.post(
                f"{API}/editions/ITS/envoyer", headers=self.entetes,
                json={**self.demande, "destinataire": "client@exemple.ci",
                      "copie": "copie@exemple.ci"},
            )
        finally:
            routeur.send_email = originel

        self.assertEqual(r.status_code, 200, r.text[:300])
        # Destinataire principal + copie
        self.assertEqual(len(appels), 2)
        for destinataire, objet, pieces in appels:
            self.assertTrue(pieces, f"aucune pièce jointe pour {destinataire}")
            nom, contenu, type_mime = pieces[0]
            self.assertTrue(nom.endswith(".pdf"))
            self.assertEqual(type_mime, "application/pdf")
            self.assertTrue(contenu.startswith(b"%PDF-"))
        self.assertEqual(r.json()["lignes"], 2)

    def test_etat_inconnu(self):
        r = client.post(f"{API}/editions/INEXISTANT/generer",
                        headers=self.entetes, json=self.demande)
        self.assertEqual(r.status_code, 404)

    def test_periode_incoherente_refusee(self):
        r = client.post(
            f"{API}/editions/JOURNAL_PAIE/generer", headers=self.entetes,
            json={**self.demande, "mois_fin": 1, "annee_fin": 2025},
        )
        self.assertEqual(r.status_code, 422)

    def test_mois_invalide_refuse(self):
        r = client.post(
            f"{API}/editions/JOURNAL_PAIE/generer", headers=self.entetes,
            json={**self.demande, "mois_debut": 13},
        )
        self.assertEqual(r.status_code, 422)

    def test_etablissement_obligatoire(self):
        """Un état est toujours rattaché à un établissement : le champ l'exige."""
        demande = {
            cle: valeur for cle, valeur in self.demande.items()
            if cle != "etablissement_id"
        }
        r = client.post(
            f"{API}/editions/JOURNAL_PAIE/generer", headers=self.entetes, json=demande,
        )
        self.assertEqual(r.status_code, 422, r.text[:200])

    def test_acces_reserve_au_cabinet(self):
        """Un état collectif expose toutes les rémunérations du dossier."""
        for role in ("client", "salarie"):
            with self.subTest(role=role):
                utilisateur = make_user(
                    self.db, f"{role}-editions@exemple.ci", role=role,
                    dossier_id=self.t["dossier"].id,
                    salarie_id=self.t["salarie"].id,
                )
                self.db.commit()
                r = client.post(
                    f"{API}/editions/JOURNAL_PAIE/generer",
                    headers=auth_headers(utilisateur), json=self.demande,
                )
                self.assertEqual(r.status_code, 403)

    def test_cloisonnement_entre_cabinets(self):
        autre = make_tenant(self.db, "cabinet-editions-autre")
        r = client.post(
            f"{API}/editions/JOURNAL_PAIE/generer",
            headers=auth_headers(autre["user"]), json=self.demande,
        )
        # Le dossier appartient à un autre cabinet
        self.assertIn(r.status_code, (403, 404))

    def test_export_refuse_a_un_role_non_habillite(self):
        utilisateur = make_user(
            self.db, "client-export@exemple.ci", role="client",
            dossier_id=self.t["dossier"].id,
        )
        self.db.commit()
        r = client.post(
            f"{API}/editions/JOURNAL_PAIE/export",
            headers=auth_headers(utilisateur), json=self.demande,
            params={"format": "pdf"},
        )
        self.assertEqual(r.status_code, 403)


if __name__ == "__main__":
    unittest.main()
